"""Information extraction: read evidence text and detect patterns in it.

Rules:
- Only report what is literally present in the text. Never fill gaps.
- `normalized` is set only when it needs no guessing (e.g. "02/03/2026" is
  ambiguous between day-first and month-first, so it stays None).
- Every entity and claim carries the source evidence ID.
"""

import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable
from urllib.parse import urlparse

from app.models import Claim, Evidence, ExtractedEntity, ExtractionResult

MAX_TEXT_CHARS = 200_000
MAX_CLAIM_CHARS = 500

# How reliably each method captures the source text (not truth of the content).
METHOD_CONFIDENCE = {"plain_text": 1.0, "python-docx": 0.95, "pypdf": 0.9, "ocr": 0.7}


# ---------------------------------------------------------------------------
# Text readers
# ---------------------------------------------------------------------------


class Unavailable(Exception):
    """Text extraction is not possible for this evidence with the current setup."""


def _read_plain(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise Unavailable("Could not decode file as UTF-8 or Windows-1252 text")


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _read_docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def run_ocr(path: Path) -> str:
    """OCR hook. Plug an engine in here; until then images report 'unavailable'."""
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        raise Unavailable("OCR engine not installed; image text was not extracted")
    try:
        return pytesseract.image_to_string(Image.open(path))
    except pytesseract.TesseractNotFoundError:
        raise Unavailable("Tesseract binary not found; image text was not extracted")


def read_text(evidence: Evidence, path: Path) -> tuple[str, str]:
    """Return (method, text) or raise Unavailable."""
    if evidence.type in ("text", "csv", "json"):
        return "plain_text", _read_plain(path)
    if evidence.type == "pdf":
        return "pypdf", _read_pdf(path)
    if evidence.type == "document":
        if path.suffix.lower() == ".docx":
            return "python-docx", _read_docx(path)
        raise Unavailable("Legacy .doc text extraction is not supported; convert to .docx or PDF")
    if evidence.type == "image":
        return "ocr", run_ocr(path)
    raise Unavailable(f"No extractor for evidence type '{evidence.type}'")


# ---------------------------------------------------------------------------
# Pattern detection
# ---------------------------------------------------------------------------

Normalizer = Callable[[str], tuple[str | None, str | None] | None]  # None = reject match

MONTH_NAMES = r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?|aug(?:ust)?|sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?)"


def _keep(raw: str) -> tuple[str | None, str | None]:
    return raw, None


def _lower(raw: str) -> tuple[str | None, str | None]:
    return raw.lower(), None


def _url(raw: str) -> tuple[str | None, str | None] | None:
    candidate = raw if raw.lower().startswith("http") else f"http://{raw}"
    parsed = urlparse(candidate)
    if not parsed.netloc or "." not in parsed.netloc:
        return None
    return raw, None


def _amount(raw: str) -> tuple[str | None, str | None] | None:
    number = re.search(r"\d(?:[\d,]*\d)?(?:\.\d{1,2})?", raw)
    if not number:
        return None
    return number.group().replace(",", ""), "INR"


def _txn(raw: str) -> tuple[str | None, str | None] | None:
    if not re.search(r"\d", raw):
        return None
    return raw.upper(), None


def _phone(raw: str) -> tuple[str | None, str | None] | None:
    digits = re.sub(r"[^\d+]", "", raw)
    if len(re.sub(r"\D", "", digits)) < 10:
        return None
    return digits, None


def _iso_datetime(raw: str) -> tuple[str | None, str | None] | None:
    for fmt in ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M"):
        try:
            return datetime.strptime(raw, fmt).isoformat(), None
        except ValueError:
            continue
    return None


def _iso_date(raw: str) -> tuple[str | None, str | None] | None:
    try:
        return datetime.strptime(raw, "%Y-%m-%d").date().isoformat(), None
    except ValueError:
        return None


def _numeric_date(raw: str) -> tuple[str | None, str | None]:
    # dd/mm vs mm/dd is ambiguous -> keep raw, do not guess.
    return None, None


def _month_date(raw: str) -> tuple[str | None, str | None]:
    cleaned = re.sub(r"(\d)(st|nd|rd|th)", r"\1", raw, flags=re.I).replace(",", "")
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    cleaned = re.sub(r"\bsept\b", "Sep", cleaned, flags=re.I).replace(".", "")
    for fmt in ("%d %B %Y", "%d %b %Y", "%B %d %Y", "%b %d %Y"):
        try:
            return datetime.strptime(cleaned, fmt).date().isoformat(), None
        except ValueError:
            continue
    return None, None


def _time(raw: str) -> tuple[str | None, str | None] | None:
    m = re.match(r"(\d{1,2})(?::(\d{2}))?(?::(\d{2}))?\s*([ap])?\.?\s*m?\.?", raw, re.I)
    if not m:
        return None
    hour, minute, second, meridiem = int(m.group(1)), int(m.group(2) or 0), m.group(3), m.group(4)
    if meridiem:
        if not 1 <= hour <= 12:
            return None
        hour = hour % 12 + (12 if meridiem.lower() == "p" else 0)
    elif hour > 23:
        return None
    if minute > 59:
        return None
    return f"{hour:02d}:{minute:02d}" + (f":{second}" if second else ""), None


# (type, compiled pattern, detection confidence, normalizer)
# Order matters: earlier patterns claim their spans first, later overlapping matches are skipped.
PATTERNS: list[tuple[str, re.Pattern[str], float, Normalizer]] = [
    ("email_address", re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}"), 0.95, _lower),
    ("url", re.compile(r"\b(?:https?://|www\.)[^\s<>\"'()\[\]{}]+", re.I), 0.95, _url),
    ("transaction_id", re.compile(r"\b(?:TXN|UTR|REF|ORD|TRX)[A-Z0-9-]*\d[A-Z0-9-]*\b", re.I), 0.9, _txn),
    (
        "transaction_id",
        re.compile(
            r"\b(?:transaction|txn|trx|reference|ref|utr|order|payment|upi\s+ref)"
            r"\s*(?:id|no|number|ref)?\.?\s*(?:[:#-]|is)?\s*(?P<v>[A-Za-z0-9][A-Za-z0-9-]{4,})",
            re.I,
        ),
        0.85,
        _txn,
    ),
    (
        "amount",
        re.compile(
            r"(?:₹|\bRs\.?|\bINR)\s?\d(?:[\d,]*\d)?(?:\.\d{1,2})?"
            r"|\b\d(?:[\d,]*\d)?(?:\.\d{1,2})?\s?(?:rupees|INR)\b",
            re.I,
        ),
        0.9,
        _amount,
    ),
    ("datetime", re.compile(r"\b\d{4}-\d{2}-\d{2}[T ]\d{2}:\d{2}(?::\d{2})?\b"), 0.9, _iso_datetime),
    ("date", re.compile(r"\b\d{4}-\d{2}-\d{2}\b"), 0.9, _iso_date),
    ("date", re.compile(r"\b\d{1,2}[/.-]\d{1,2}[/.-]\d{2,4}\b"), 0.8, _numeric_date),
    (
        "date",
        re.compile(
            rf"\b\d{{1,2}}(?:st|nd|rd|th)?\s+{MONTH_NAMES}\.?,?\s+\d{{4}}\b"
            rf"|\b{MONTH_NAMES}\.?\s+\d{{1,2}}(?:st|nd|rd|th)?,?\s+\d{{4}}\b",
            re.I,
        ),
        0.85,
        _month_date,
    ),
    (
        "time",
        re.compile(r"\b\d{1,2}:\d{2}(?::\d{2})?(?:\s*(?:[ap]m\b|[ap]\.m\.))?|\b\d{1,2}\s*(?:[ap]m\b|[ap]\.m\.)", re.I),
        0.85,
        _time,
    ),
    (
        "phone_number",
        re.compile(r"(?<![\w+])(?:\+91[\s-]?|0)?[6-9]\d{4}[\s-]?\d{5}(?!\d)|(?<![\w+])\+\d{1,3}[\s-]?\d[\d\s-]{7,13}\d(?!\d)"),
        0.8,
        _phone,
    ),
]

TRAILING_PUNCT = ".,;:!?)]}'\""


def find_entities(text: str, evidence_id: str) -> list[ExtractedEntity]:
    taken: list[tuple[int, int]] = []
    seen: set[tuple[str, str]] = set()
    found: list[tuple[int, int, str, str, str | None, str | None, float]] = []

    for etype, pattern, confidence, normalize in PATTERNS:
        for m in pattern.finditer(text):
            start, end = m.span("v") if "v" in pattern.groupindex else m.span()
            if etype == "url":
                while end > start and text[end - 1] in TRAILING_PUNCT:
                    end -= 1
            if any(start < t_end and t_start < end for t_start, t_end in taken):
                continue
            raw = text[start:end]
            result = normalize(raw)
            if result is None:
                continue
            taken.append((start, end))
            normalized, unit = result
            key = (etype, (normalized or raw).lower())
            if key in seen:
                continue  # same value repeated in the text; report once
            seen.add(key)
            found.append((start, end, etype, raw, normalized, unit, confidence))

    found.sort(key=lambda f: f[0])
    return [
        ExtractedEntity(
            id=f"{evidence_id}-E{i:02d}",
            type=etype,
            raw=raw,
            normalized=normalized,
            unit=unit,
            source_evidence_ids=[evidence_id],
            confidence=confidence,
            char_start=start,
            char_end=end,
        )
        for i, (start, end, etype, raw, normalized, unit, confidence) in enumerate(found, 1)
    ]


# ---------------------------------------------------------------------------
# Claims: sentences / lines that contain at least one detected entity, verbatim.
# ---------------------------------------------------------------------------

ABBREVIATIONS = {"rs", "no", "mr", "mrs", "ms", "dr", "a.m", "p.m", "ref", "st", "vs", "e.g", "i.e"}


def _sentence_spans(text: str) -> list[tuple[int, int]]:
    spans: list[tuple[int, int]] = []
    start = 0
    for m in re.finditer(r"\n+|[.!?]+(?=\s|$)", text):
        if not m.group().startswith("\n"):
            prev = re.search(r"([\w.]+)$", text[start : m.start()])
            if prev and prev.group(1).lower() in ABBREVIATIONS:
                continue
        end = m.end() if not m.group().startswith("\n") else m.start()
        spans.append((start, end))
        start = m.end()
    spans.append((start, len(text)))

    trimmed = []
    for s, e in spans:
        while s < e and text[s].isspace():
            s += 1
        while e > s and text[e - 1].isspace():
            e -= 1
        if e > s:
            trimmed.append((s, e))
    return trimmed


def build_claims(text: str, entities: list[ExtractedEntity], evidence_id: str, confidence: float) -> list[Claim]:
    claims = []
    for s, e in _sentence_spans(text):
        inside = [ent.id for ent in entities if s <= ent.char_start and ent.char_end <= e]
        if not inside:
            continue
        claims.append(
            Claim(
                id="",  # assigned by the extraction store
                claim=text[s:e][:MAX_CLAIM_CHARS],
                source_evidence_ids=[evidence_id],
                confidence=confidence,
                entity_ids=inside,
                char_start=s,
                char_end=e,
            )
        )
    return claims


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def extract(evidence: Evidence, path: Path) -> ExtractionResult:
    now = datetime.now(timezone.utc)
    base = {"evidence_id": evidence.id, "extracted_at": now}

    try:
        method, text = read_text(evidence, path)
    except Unavailable as e:
        return ExtractionResult(**base, status="unavailable", notes=[str(e)])
    except Exception as e:  # corrupt/unreadable file
        return ExtractionResult(**base, status="failed", notes=[f"Text extraction failed: {e}"])

    notes: list[str] = []
    truncated = len(text) > MAX_TEXT_CHARS
    if truncated:
        text = text[:MAX_TEXT_CHARS]
        notes.append(f"Text truncated to first {MAX_TEXT_CHARS:,} characters")

    if not text.strip():
        if evidence.type == "pdf":
            notes.append("PDF has no text layer; it may be scanned and require OCR")
        else:
            notes.append("No text content found")
        return ExtractionResult(
            **base, status="no_text", method=method, extracted_text=text, text_truncated=truncated, notes=notes
        )

    confidence = METHOD_CONFIDENCE[method]
    entities = find_entities(text, evidence.id)
    claims = build_claims(text, entities, evidence.id, confidence)
    return ExtractionResult(
        **base,
        status="extracted",
        method=method,
        extracted_text=text,
        text_truncated=truncated,
        entities=entities,
        claims=claims,
        confidence=confidence,
        notes=notes,
    )
